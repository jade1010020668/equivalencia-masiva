"""
Constructor del Excel consolidado — columnas alineadas al notebook mismo_grado_.ipynb.

El reporte contiene una fila por cada par (base 4.0, lista candidata) con
todos los scores y razones que produce el notebook, más metadatos de auditoría.
"""

from __future__ import annotations

import io
from collections import Counter
from dataclasses import asdict, dataclass, field

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.decision import DecisionFinal
from core.semantic import AspectosResult


# =========================================================================
# Estructura del reporte
# =========================================================================


@dataclass
class FilaReporte:
    """Una fila del Excel consolidado — refleja un par (base, lista) analizado."""

    # Identificación
    entidad: str
    grupo: str

    # Base 4.0
    opec_base: str
    codigo_base: str
    denom_base: str
    grado_base: str
    nivel_base: str
    salario_base: float

    # Lista candidata
    opec_lista: str
    codigo_lista: str
    denom_lista: str
    grado_lista: str
    nivel_lista: str
    salario_lista: float

    # Decisión final (notebook: es_equivalente + equivalence_score)
    decision_final: str  # EQUIVALENTE / NO_EQUIVALENTE
    es_equivalente: bool
    score_ponderado: float
    razones_no_eq: str  # razones unidas por " | "

    # Gating
    nivel_pasa: bool
    nivel_msg: str
    salario_pasa: bool
    salario_msg: str

    # Coberturas SBERT (cobertura_prom_max_1 del notebook)
    sbert_funciones: float
    sbert_estudio: float
    sbert_comp_lab: float
    sbert_comp_comp: float
    sbert_experiencia: float

    # Pasa / No pasa por aspecto (según umbrales)
    pasa_funciones: bool
    pasa_estudio: bool
    pasa_comp_lab: bool
    pasa_comp_comp: bool

    # Experiencia y jerarquía
    tipo_exp_base: str
    tipo_exp_lista: str
    exp_jerarquia_msg: str

    # Modelo usado
    modelo_sbert: str

    # --- Textos originales para auditoría ---
    funciones_base: str = ""
    funciones_lista: str = ""
    estudio_base: str = ""
    estudio_lista: str = ""
    experiencia_base: str = ""
    experiencia_lista: str = ""
    comp_lab_base: str = ""
    comp_lab_lista: str = ""
    comp_comp_base: str = ""
    comp_comp_lista: str = ""


# =========================================================================
# Construcción de filas
# =========================================================================


def _read(row: pd.Series, col: str, default="") -> str:
    if col not in row.index:
        return default
    val = row[col]
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return default
    return val


def construir_fila(
    grupo_nombre: str,
    entidad: str,
    base: pd.Series,
    lista: pd.Series,
    aspectos: AspectosResult,
    decision: DecisionFinal,
) -> FilaReporte:
    """Arma la fila combinando Series de entrada + resultado semántico + decisión."""
    return FilaReporte(
        entidad=str(entidad),
        grupo=grupo_nombre,
        opec_base=str(_read(base, "No. OPEC")),
        codigo_base=str(_read(base, "Código")),
        denom_base=str(_read(base, "Denominación")),
        grado_base=str(_read(base, "Grado")),
        nivel_base=str(_read(base, "Nivel Jerárquico")),
        salario_base=float(_read(base, "Salario", 0) or 0),
        opec_lista=str(_read(lista, "No. OPEC")),
        codigo_lista=str(_read(lista, "Código")),
        denom_lista=str(_read(lista, "Denominación")),
        grado_lista=str(_read(lista, "Grado")),
        nivel_lista=str(_read(lista, "Nivel Jerárquico")),
        salario_lista=float(_read(lista, "Salario", 0) or 0),
        decision_final=decision.clasificacion,
        es_equivalente=decision.es_equivalente,
        score_ponderado=decision.score,
        razones_no_eq=" | ".join(decision.razones) if decision.razones else "",
        nivel_pasa=aspectos.nivel_pasa,
        nivel_msg=aspectos.nivel_msg,
        salario_pasa=aspectos.salario_pasa,
        salario_msg=aspectos.salario_msg,
        sbert_funciones=aspectos.funciones.cobertura,
        sbert_estudio=aspectos.estudio.cobertura,
        sbert_comp_lab=aspectos.comp_lab.cobertura,
        sbert_comp_comp=aspectos.comp_comp.cobertura,
        sbert_experiencia=aspectos.experiencia.cobertura,
        pasa_funciones=decision.pasa_funciones,
        pasa_estudio=decision.pasa_estudio,
        pasa_comp_lab=decision.pasa_comp_lab,
        pasa_comp_comp=decision.pasa_comp_comp,
        tipo_exp_base=aspectos.tipo_exp_base.value,
        tipo_exp_lista=aspectos.tipo_exp_destino.value,
        exp_jerarquia_msg=aspectos.experiencia.jerarquia_msg or "",
        modelo_sbert=aspectos.modelo_usado,
        funciones_base=str(_read(base, "Funciones")),
        funciones_lista=str(_read(lista, "Funciones")),
        estudio_base=str(_read(base, "Requisitos Estudio")),
        estudio_lista=str(_read(lista, "Requisitos Estudio")),
        experiencia_base=str(_read(base, "Requisitos Experiencia")),
        experiencia_lista=str(_read(lista, "Requisitos Experiencia")),
        comp_lab_base=str(_read(base, "Competencias Laborales")),
        comp_lab_lista=str(_read(lista, "Competencias Laborales")),
        comp_comp_base=str(_read(base, "Competencias Comportamentales")),
        comp_comp_lista=str(_read(lista, "Competencias Comportamentales")),
    )


# =========================================================================
# Exportador Excel
# =========================================================================

_HEADERS: dict[str, str] = {
    "entidad": "Entidad",
    "grupo": "Grupo",
    "opec_base": "OPEC Base",
    "codigo_base": "Código Base",
    "denom_base": "Denominación Base",
    "grado_base": "Grado Base",
    "nivel_base": "Nivel Base",
    "salario_base": "Salario Base",
    "opec_lista": "OPEC Lista",
    "codigo_lista": "Código Lista",
    "denom_lista": "Denominación Lista",
    "grado_lista": "Grado Lista",
    "nivel_lista": "Nivel Lista",
    "salario_lista": "Salario Lista",
    "decision_final": "Decisión Final",
    "es_equivalente": "¿Es Equivalente?",
    "score_ponderado": "Score Ponderado",
    "razones_no_eq": "Razones No Equivalencia",
    "nivel_pasa": "Nivel Pasa",
    "nivel_msg": "Nivel (detalle)",
    "salario_pasa": "Salario Pasa",
    "salario_msg": "Salario (detalle)",
    "sbert_funciones": "SBERT Funciones",
    "sbert_estudio": "SBERT Estudio",
    "sbert_comp_lab": "SBERT Comp. Laborales",
    "sbert_comp_comp": "SBERT Comp. Comport.",
    "sbert_experiencia": "SBERT Experiencia",
    "pasa_funciones": "Pasa Funciones",
    "pasa_estudio": "Pasa Estudio",
    "pasa_comp_lab": "Pasa Comp. Lab.",
    "pasa_comp_comp": "Pasa Comp. Comp.",
    "tipo_exp_base": "Tipo Exp. Base",
    "tipo_exp_lista": "Tipo Exp. Lista",
    "exp_jerarquia_msg": "Jerarquía Experiencia",
    "modelo_sbert": "Modelo SBERT",
    "funciones_base": "Funciones Base (texto)",
    "funciones_lista": "Funciones Lista (texto)",
    "estudio_base": "Estudio Base (texto)",
    "estudio_lista": "Estudio Lista (texto)",
    "experiencia_base": "Experiencia Base (texto)",
    "experiencia_lista": "Experiencia Lista (texto)",
    "comp_lab_base": "Comp. Lab. Base (texto)",
    "comp_lab_lista": "Comp. Lab. Lista (texto)",
    "comp_comp_base": "Comp. Comp. Base (texto)",
    "comp_comp_lista": "Comp. Comp. Lista (texto)",
}

_PERCENT_FIELDS = {
    "score_ponderado",
    "sbert_funciones", "sbert_estudio", "sbert_comp_lab", "sbert_comp_comp", "sbert_experiencia",
}
_CURRENCY_FIELDS = {"salario_base", "salario_lista"}
_WRAP_FIELDS = {
    "razones_no_eq", "nivel_msg", "salario_msg", "exp_jerarquia_msg",
    "funciones_base", "funciones_lista", "estudio_base", "estudio_lista",
    "experiencia_base", "experiencia_lista",
    "comp_lab_base", "comp_lab_lista", "comp_comp_base", "comp_comp_lista",
}

_COLORS: dict[str, str] = {
    "EQUIVALENTE": "C6EFCE",       # verde claro
    "NO_EQUIVALENTE": "FFC7CE",    # rojo claro
}


def filas_a_dataframe(filas: list[FilaReporte]) -> pd.DataFrame:
    if not filas:
        return pd.DataFrame(columns=list(_HEADERS.keys()))
    return pd.DataFrame([asdict(f) for f in filas])


def exportar_excel(filas: list[FilaReporte], nombre_archivo: str = "reporte") -> bytes:
    """Genera un .xlsx con las filas + pestaña Resumen + formato condicional."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Análisis Masivo"

    campos = list(_HEADERS.keys())
    labels = [_HEADERS[c] for c in campos]

    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for idx, label in enumerate(labels, start=1):
        cell = ws.cell(row=1, column=idx, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "C2"

    for r, fila in enumerate(filas, start=2):
        d = asdict(fila)
        for c, campo in enumerate(campos, start=1):
            valor = d[campo]
            cell = ws.cell(row=r, column=c, value=valor)
            if campo in _PERCENT_FIELDS and isinstance(valor, (int, float)):
                cell.number_format = "0.0%"
            elif campo in _CURRENCY_FIELDS and isinstance(valor, (int, float)):
                cell.number_format = "#,##0"
            elif campo in _WRAP_FIELDS:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        # Color por clasificación
        color = _COLORS.get(fila.decision_final)
        if color:
            dec_col = campos.index("decision_final") + 1
            ws.cell(row=r, column=dec_col).fill = PatternFill(
                start_color=color, end_color=color, fill_type="solid"
            )

    col_widths: dict[str, int] = {
        "entidad": 28, "grupo": 38,
        "opec_base": 12, "codigo_base": 12, "denom_base": 32,
        "grado_base": 10, "nivel_base": 16, "salario_base": 14,
        "opec_lista": 12, "codigo_lista": 12, "denom_lista": 32,
        "grado_lista": 10, "nivel_lista": 16, "salario_lista": 14,
        "decision_final": 18, "es_equivalente": 12, "score_ponderado": 14,
        "razones_no_eq": 50,
        "nivel_pasa": 10, "nivel_msg": 32, "salario_pasa": 10, "salario_msg": 32,
        "sbert_funciones": 14, "sbert_estudio": 14, "sbert_comp_lab": 14,
        "sbert_comp_comp": 14, "sbert_experiencia": 14,
        "pasa_funciones": 10, "pasa_estudio": 10, "pasa_comp_lab": 10, "pasa_comp_comp": 10,
        "tipo_exp_base": 22, "tipo_exp_lista": 22, "exp_jerarquia_msg": 40,
        "modelo_sbert": 38,
        "funciones_base": 60, "funciones_lista": 60,
        "estudio_base": 40, "estudio_lista": 40,
        "experiencia_base": 40, "experiencia_lista": 40,
        "comp_lab_base": 40, "comp_lab_lista": 40,
        "comp_comp_base": 40, "comp_comp_lista": 40,
    }
    for idx, campo in enumerate(campos, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = col_widths.get(campo, 14)

    # Pestaña Resumen
    resumen = wb.create_sheet("Resumen")
    resumen.append(["Clasificación", "Pares"])
    counter = Counter(f.decision_final for f in filas)
    for clave in ["EQUIVALENTE", "NO_EQUIVALENTE"]:
        resumen.append([clave, counter.get(clave, 0)])
    resumen.append(["TOTAL", len(filas)])
    resumen.column_dimensions["A"].width = 22
    resumen.column_dimensions["B"].width = 12
    for cell in resumen["1:1"]:
        cell.fill = header_fill
        cell.font = header_font

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
