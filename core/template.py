"""
Generador de plantilla Excel descargable.

Crea un .xlsx con TODAS las columnas que la app espera, con encabezados
estilizados y 2-3 filas de ejemplo (un empleo base 4.0 + sus listas candidatas)
para que el usuario entienda el formato.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import REQUIRED_COLUMNS


# Ayuda contextual para cada columna (tooltips en el encabezado)
_COLUMN_HELP: dict[str, str] = {
    "No. OPEC": "Número de la OPEC (Oferta Pública de Empleo de Carrera). Identificador único del empleo.",
    "Código": "Código del empleo según el manual de funciones de la entidad.",
    "Denominación": "Nombre del cargo, por ejemplo 'Profesional Universitario'.",
    "Grado": "Grado del empleo (número entero).",
    "Nivel Jerárquico": "Directivo / Asesor / Profesional / Técnico / Asistencial.",
    "Orden": "Nacional / Territorial.",
    "Naturaleza Jurídica": "Tipo de entidad: Ministerio, Departamento, Municipio, etc.",
    "Salario": "Salario mensual del cargo en pesos colombianos (número).",
    "Requisitos Estudio": "Requisitos académicos textuales del manual de funciones.",
    "Requisitos Experiencia": "Experiencia laboral/profesional requerida (con tipo).",
    "Funciones": "Funciones del cargo, preferiblemente separadas por viñetas o saltos de línea.",
    "Competencias Laborales": "Competencias laborales exigidas para el cargo.",
    "Competencias Comportamentales": "Competencias comportamentales exigidas.",
    "modalidad": "'4.0' para empleo base vigente, 'LISTAS' para cada lista de elegibles candidata.",
    "nombre_entidad": "Nombre de la entidad titular del empleo.",
}


# Filas de ejemplo (un base 4.0 + dos listas candidatas de la misma entidad/nivel/grado)
_EJEMPLOS = [
    {
        "No. OPEC": 999001,
        "Código": 219,
        "Denominación": "Profesional Universitario",
        "Grado": 7,
        "Nivel Jerárquico": "Profesional",
        "Orden": "Territorial",
        "Naturaleza Jurídica": "Municipio",
        "Salario": 3200000,
        "Requisitos Estudio": (
            "Título profesional en Derecho o en Ciencias Políticas "
            "y Tarjeta profesional vigente en los casos reglamentados por la ley."
        ),
        "Requisitos Experiencia": "Experiencia profesional relacionada de veinticuatro (24) meses.",
        "Funciones": (
            "- Asesorar jurídicamente al despacho en temas normativos y procesales.\n"
            "- Elaborar conceptos jurídicos sobre actos administrativos.\n"
            "- Representar legalmente a la entidad en procesos judiciales.\n"
            "- Las demás funciones asignadas por el jefe inmediato."
        ),
        "Competencias Laborales": (
            "- Orientación a resultados.\n"
            "- Compromiso con la organización.\n"
            "- Trabajo en equipo."
        ),
        "Competencias Comportamentales": (
            "- Aprendizaje continuo.\n"
            "- Experticia profesional.\n"
            "- Trabajo en equipo y colaboración."
        ),
        "modalidad": "4.0",
        "nombre_entidad": "Alcaldía de Cali (EJEMPLO)",
    },
    {
        "No. OPEC": 999002,
        "Código": 219,
        "Denominación": "Profesional Universitario",
        "Grado": 7,
        "Nivel Jerárquico": "Profesional",
        "Orden": "Territorial",
        "Naturaleza Jurídica": "Municipio",
        "Salario": 3200000,
        "Requisitos Estudio": (
            "Título profesional en Derecho o en Ciencias Políticas."
        ),
        "Requisitos Experiencia": "Veinticuatro (24) meses de experiencia profesional relacionada.",
        "Funciones": (
            "- Asesorar jurídicamente a la oficina en temas normativos y procedimentales.\n"
            "- Elaborar conceptos legales sobre actos administrativos.\n"
            "- Representar a la entidad ante autoridades judiciales."
        ),
        "Competencias Laborales": (
            "- Orientación al logro.\n"
            "- Trabajo en equipo."
        ),
        "Competencias Comportamentales": (
            "- Aprendizaje continuo.\n"
            "- Experticia profesional."
        ),
        "modalidad": "LISTAS",
        "nombre_entidad": "Alcaldía de Cali (EJEMPLO)",
    },
    {
        "No. OPEC": 999003,
        "Código": 219,
        "Denominación": "Profesional Universitario",
        "Grado": 7,
        "Nivel Jerárquico": "Profesional",
        "Orden": "Territorial",
        "Naturaleza Jurídica": "Municipio",
        "Salario": 3200000,
        "Requisitos Estudio": (
            "Título profesional en Administración Pública o en Ciencias Sociales."
        ),
        "Requisitos Experiencia": "Dieciocho (18) meses de experiencia laboral relacionada.",
        "Funciones": (
            "- Gestionar proyectos sociales con comunidades vulnerables.\n"
            "- Coordinar actividades de formación ciudadana.\n"
            "- Preparar informes trimestrales de avance."
        ),
        "Competencias Laborales": (
            "- Orientación al usuario y al ciudadano.\n"
            "- Capacidad de análisis."
        ),
        "Competencias Comportamentales": (
            "- Creatividad e innovación.\n"
            "- Trabajo en equipo."
        ),
        "modalidad": "LISTAS",
        "nombre_entidad": "Alcaldía de Cali (EJEMPLO)",
    },
]


def generar_plantilla_excel() -> bytes:
    """
    Crea el archivo .xlsx de plantilla en memoria y devuelve los bytes.
    Incluye:
      - Encabezados estilizados con tooltips (openpyxl Comments).
      - 3 filas de ejemplo con estructura válida (1 base 4.0 + 2 listas).
      - Una segunda pestaña 'Instrucciones' con explicación de cada columna.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "OPEC"

    header_fill = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)

    # Encabezado
    for idx, col in enumerate(REQUIRED_COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx, value=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        if col in _COLUMN_HELP:
            cell.comment = Comment(_COLUMN_HELP[col], "Plantilla Despacho EARM")

    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "A2"

    # Filas de ejemplo
    for r, ejemplo in enumerate(_EJEMPLOS, start=2):
        for c, col in enumerate(REQUIRED_COLUMNS, start=1):
            cell = ws.cell(row=r, column=c, value=ejemplo.get(col, ""))
            if col in {"Funciones", "Requisitos Estudio", "Requisitos Experiencia",
                       "Competencias Laborales", "Competencias Comportamentales"}:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col == "Salario":
                cell.number_format = "#,##0"

    # Anchos de columnas
    anchos = {
        "No. OPEC": 12, "Código": 10, "Denominación": 28,
        "Grado": 8, "Nivel Jerárquico": 18, "Orden": 14,
        "Naturaleza Jurídica": 20, "Salario": 14,
        "Requisitos Estudio": 40, "Requisitos Experiencia": 32,
        "Funciones": 55, "Competencias Laborales": 35,
        "Competencias Comportamentales": 35,
        "modalidad": 12, "nombre_entidad": 28,
    }
    for idx, col in enumerate(REQUIRED_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = anchos.get(col, 18)

    # Pestaña de instrucciones
    instr = wb.create_sheet("Instrucciones")
    instr.append(["Columna", "Qué poner aquí"])
    for cell in instr["1:1"]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    for col in REQUIRED_COLUMNS:
        instr.append([col, _COLUMN_HELP.get(col, "")])
    instr.column_dimensions["A"].width = 28
    instr.column_dimensions["B"].width = 90
    for row in instr.iter_rows(min_row=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    nota_final = (
        "NOTAS IMPORTANTES:\n"
        "1. Cada empleo base VIGENTE debe tener modalidad='4.0'.\n"
        "2. Cada lista de elegibles candidata debe tener modalidad='LISTAS'.\n"
        "3. Las listas candidatas DEBEN tener el MISMO 'Grado' que el empleo base para ser agrupadas.\n"
        "4. 'nombre_entidad' y 'Nivel Jerárquico' deben coincidir entre el base y sus listas.\n"
        "5. Las filas de ejemplo traen OPEC 999001-999003 — borra o reemplaza antes de usar.\n"
        "6. Los campos largos (Funciones, Estudio, etc.) admiten viñetas (- o •) y saltos de línea."
    )
    instr.append([])
    for linea in nota_final.splitlines():
        instr.append(["", linea])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()
